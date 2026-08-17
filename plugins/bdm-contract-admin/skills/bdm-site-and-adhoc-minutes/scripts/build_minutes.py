#!/usr/bin/env python3
"""build_minutes.py — populate a Form 232 Meeting Memorandum template from a
JSON config and produce a new .docx.

Handles the BDM template's quirks:
  - Meta table cells.
  - Attendees/Apologies tables with auto row trimming or expansion.
  - Topic section paragraphs with TWO-RUN formatting:
        leading "NN " in tan (9D7B5B), title in dark navy (0F1721) bold
        leading "•\t" in tan, body in dark navy
  - Action register table with auto row expansion.
  - Distribution + Next Meeting + Summary paragraphs.

The output still needs the `bdm-pdf-export` preflight before PDF conversion.

Usage:
    python3 build_minutes.py \\
        --template path/to/232-Meeting_Memorandum_R2_2026-05.docx \\
        --config path/to/minutes_config.json \\
        --output path/to/<prefix>###_<project>.docx
"""

from __future__ import annotations

import argparse
import json
import sys
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn


# Two-run colour scheme used in BDM templates' topic paragraphs
PREFIX_COLOUR = "9D7B5B"
BODY_COLOUR = "0F1721"


def two_run_set(para, text: str) -> None:
    """Split text into (prefix, body) and write into first two runs.
    Prefix detected as either '•\\t' (bullet) or 'NN ' (two digits + space)."""
    if text.startswith("•\t"):
        prefix, body = "•\t", text[2:]
    elif len(text) >= 3 and text[0:2].isdigit() and text[2] == " ":
        prefix, body = text[0:3], text[3:]
    else:
        prefix, body = "", text

    runs = para.runs
    if len(runs) >= 2 and prefix:
        runs[0].text = prefix
        runs[1].text = body
        for r in runs[2:]:
            if r.text:
                r.text = ""
    elif runs:
        runs[0].text = text
        for r in runs[1:]:
            if r.text:
                r.text = ""
    else:
        para.add_run(text)


def set_cell(cell, text: str) -> None:
    paras = cell.paragraphs
    if not paras:
        cell.add_paragraph(text)
        return
    runs = paras[0].runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            if r.text:
                r.text = ""
    else:
        paras[0].add_run(text)
    for p in paras[1:]:
        for r in p.runs:
            r.text = ""


def find_paragraph_by_prefix(doc, prefix: str):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None


def populate_meta_table(table, meta: dict[str, str]) -> None:
    rows = table.rows
    set_cell(rows[0].cells[1], meta.get("project", ""))
    set_cell(rows[1].cells[1], meta.get("meeting_no", ""))
    set_cell(rows[2].cells[1], meta.get("location", ""))
    set_cell(rows[3].cells[1], meta.get("date_time", ""))


def populate_attendees_table(table, attendees: list[dict], apologies: list[dict]) -> None:
    """Template has 6 rows: header, 3 attendee slots, APOLOGIES separator,
    1 apology slot. Trim or expand to fit actual lists."""
    HEADER_ROW = 0
    SEPARATOR_TEXT = "APOLOGIES"

    # Find separator row index (any row where all cells say APOLOGIES)
    sep_idx = None
    for i, row in enumerate(table.rows):
        if all(c.text.strip().upper() == SEPARATOR_TEXT for c in row.cells):
            sep_idx = i
            break
    if sep_idx is None:
        # Fall back to expected layout
        sep_idx = 4

    # Attendee slot rows: between header and separator (exclusive of both)
    attendee_slot_rows = list(table.rows)[HEADER_ROW + 1: sep_idx]
    # Apology slot rows: after separator
    apology_slot_rows = list(table.rows)[sep_idx + 1:]

    def set_row(row, data: dict) -> None:
        cells = row.cells
        set_cell(cells[0], data.get("name", ""))
        set_cell(cells[1], data.get("initials", ""))
        set_cell(cells[2], data.get("company", ""))
        set_cell(cells[3], data.get("email", ""))

    # Populate / remove attendee slots
    for i, person in enumerate(attendees):
        if i < len(attendee_slot_rows):
            set_row(attendee_slot_rows[i], person)
        else:
            # Clone last filled slot row
            template_row = (
                attendee_slot_rows[-1] if attendee_slot_rows
                else table.rows[1]
            )
            new_tr = deepcopy(template_row._tr)
            template_row._tr.addnext(new_tr)
            attendee_slot_rows.append(table.rows[attendee_slot_rows[-1]._tr.getparent().index(new_tr)])
            # Easier: re-fetch by walking
            new_rows = list(table.rows)
            target = new_rows[HEADER_ROW + 1 + i]
            set_row(target, person)

    # Remove unused attendee slots
    used = len(attendees)
    for slot_row in attendee_slot_rows[used:]:
        tr = slot_row._tr
        tr.getparent().remove(tr)

    # Re-fetch separator/apology rows after attendee mutations
    rows = list(table.rows)
    # Find separator again
    new_sep_idx = None
    for i, row in enumerate(rows):
        if all(c.text.strip().upper() == SEPARATOR_TEXT for c in row.cells):
            new_sep_idx = i
            break
    if new_sep_idx is None:
        return
    apology_slot_rows = rows[new_sep_idx + 1:]

    for i, person in enumerate(apologies):
        if i < len(apology_slot_rows):
            set_row(apology_slot_rows[i], person)
        else:
            template_row = (
                apology_slot_rows[-1] if apology_slot_rows
                else rows[new_sep_idx]
            )
            new_tr = deepcopy(template_row._tr)
            template_row._tr.addnext(new_tr)
            new_rows = list(table.rows)
            target = new_rows[new_sep_idx + 1 + i]
            set_row(target, person)

    # Remove unused apology slots
    rows = list(table.rows)
    new_sep_idx = None
    for i, row in enumerate(rows):
        if all(c.text.strip().upper() == SEPARATOR_TEXT for c in row.cells):
            new_sep_idx = i
            break
    apology_slot_rows = rows[new_sep_idx + 1:]
    for slot_row in apology_slot_rows[len(apologies):]:
        tr = slot_row._tr
        tr.getparent().remove(tr)


def populate_topics(doc, topics: list[dict]) -> None:
    """Replace the template's 3 placeholder topic sections with N sections from config.
    Each topic = {heading: '01 ...', bullets: [...]}.

    Strategy: clone the existing first topic-header paragraph and first bullet
    paragraph as templates, wipe the placeholder section paragraphs between
    'TOPICS DISCUSSED' and 'ACTION REGISTER', then insert new content using
    those clones.
    """
    paras = list(doc.paragraphs)
    topics_header_p = None
    action_register_p = None
    for p in paras:
        t = p.text.strip()
        if t == "TOPICS DISCUSSED" and topics_header_p is None:
            topics_header_p = p
        elif t == "ACTION REGISTER" and action_register_p is None:
            action_register_p = p

    if topics_header_p is None or action_register_p is None:
        raise RuntimeError("Could not find TOPICS DISCUSSED / ACTION REGISTER anchors")

    # Find first topic header paragraph (e.g. '01 Topic Heading One')
    th_idx = paras.index(topics_header_p)
    ar_idx = paras.index(action_register_p)
    placeholder_paras = paras[th_idx + 1:ar_idx]

    # Clone XML of first topic header and first bullet for re-use
    header_xml = None
    bullet_xml = None
    for p in placeholder_paras:
        text = p.text.strip()
        if not text:
            continue
        if (len(text) >= 3 and text[0:2].isdigit() and text[2].isspace()) and header_xml is None:
            header_xml = deepcopy(p._element)
        elif text.startswith("•") and bullet_xml is None:
            bullet_xml = deepcopy(p._element)
        if header_xml is not None and bullet_xml is not None:
            break

    if header_xml is None or bullet_xml is None:
        raise RuntimeError("Could not find placeholder topic header / bullet in template")

    # Remove all placeholder topic paragraphs
    ar_elem = action_register_p._element
    parent = ar_elem.getparent()
    e = topics_header_p._element.getnext()
    while e is not None and e is not ar_elem:
        nxt = e.getnext()
        parent.remove(e)
        e = nxt

    # Insert new content
    def set_text_in_xml(p_xml, text: str) -> None:
        if text.startswith("•\t"):
            prefix, body = "•\t", text[2:]
        elif len(text) >= 3 and text[0:2].isdigit() and text[2] == " ":
            prefix, body = text[0:3], text[3:]
        else:
            prefix, body = "", text
        t_elems = p_xml.findall(".//" + qn("w:t"))
        if len(t_elems) >= 2 and prefix:
            t_elems[0].text = prefix
            t_elems[1].text = body
            for t in t_elems[2:]:
                t.text = ""
        elif t_elems:
            t_elems[0].text = text
            for t in t_elems[1:]:
                t.text = ""

    for topic in topics:
        h_xml = deepcopy(header_xml)
        set_text_in_xml(h_xml, topic["heading"])
        ar_elem.addprevious(h_xml)
        for bullet in topic.get("bullets", []):
            b_xml = deepcopy(bullet_xml)
            set_text_in_xml(b_xml, "•\t" + bullet)
            ar_elem.addprevious(b_xml)


def populate_action_register(table, actions: list[dict]) -> None:
    """Update the template's placeholder action rows (typically 4) and grow as needed."""
    data_rows = list(table.rows[1:])

    def set_row(row, action: dict) -> None:
        set_cell(row.cells[0], action.get("owner", ""))
        set_cell(row.cells[1], action.get("action", ""))
        set_cell(row.cells[2], action.get("due", ""))

    # Update existing rows
    for i in range(min(len(data_rows), len(actions))):
        set_row(data_rows[i], actions[i])

    # Add extra rows
    for i in range(len(data_rows), len(actions)):
        last_tr = table.rows[-1]._tr
        new_tr = deepcopy(last_tr)
        last_tr.addnext(new_tr)
        new_row = table.rows[-1]
        set_row(new_row, actions[i])

    # Remove unused placeholder rows
    if len(actions) < len(data_rows):
        rows = list(table.rows)
        for slot_row in rows[1 + len(actions):]:
            tr = slot_row._tr
            tr.getparent().remove(tr)


def populate_distribution(doc, distribution: str) -> None:
    if not distribution:
        return
    for p in doc.paragraphs:
        if p.text.strip().startswith("DISTRIBUTION"):
            for r in p.runs:
                if "[additional distribution list]" in r.text:
                    r.text = r.text.replace(
                        "[additional distribution list]", distribution
                    )
            return


def populate_summary(doc, summary: str) -> None:
    if not summary:
        return
    for p in doc.paragraphs:
        if p.text.strip().startswith("Provide a short narrative summary"):
            two_run_set(p, summary)
            return
    # Fallback: SUMMARY anchor
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "SUMMARY" and i + 1 < len(doc.paragraphs):
            two_run_set(doc.paragraphs[i + 1], summary)
            return


def populate_next_meeting(doc, text: str) -> None:
    if not text:
        return
    for p in doc.paragraphs:
        if "[Day, DD Month YYYY at HH:MM" in p.text:
            for r in p.runs:
                r.text = ""
            if p.runs:
                p.runs[0].text = text
            return


def build(template: Path, config: dict, output: Path) -> None:
    doc = Document(str(template))

    # Meta
    populate_meta_table(doc.tables[0], config.get("meta", {}))

    # Attendees + apologies
    populate_attendees_table(
        doc.tables[1],
        config.get("attendees", []),
        config.get("apologies", []),
    )

    # Distribution
    populate_distribution(doc, config.get("distribution", ""))

    # Summary
    populate_summary(doc, config.get("summary", ""))

    # Topics
    populate_topics(doc, config.get("topics", []))

    # Action register
    populate_action_register(doc.tables[2], config.get("actions", []))

    # Next meeting
    populate_next_meeting(doc, config.get("next_meeting", ""))

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    print(f"Saved: {output}")


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True, type=Path)
    ap.add_argument("--config", required=True, type=Path)
    ap.add_argument("--output", required=True, type=Path)
    args = ap.parse_args(argv[1:])

    cfg = json.loads(args.config.read_text(encoding="utf-8"))
    build(args.template, cfg, args.output)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
