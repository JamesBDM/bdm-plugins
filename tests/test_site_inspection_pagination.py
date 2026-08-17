import base64
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


REPO_ROOT = Path(__file__).resolve().parents[1]
BUILDER = (
    REPO_ROOT
    / "plugins"
    / "bdm-contract-admin"
    / "skills"
    / "bdm-site-inspection-report"
    / "scripts"
    / "build_report.py"
)

# Synthetic portrait PNG payload. The regression was discovered on a ten-photo site
# report, but no project data or photographs belong in the public repository.
PORTRAIT_IMAGE = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAABQAAAAeCAIAAACjcKk8AAAANUlEQVR4nO3LoQEAIAwDwZD9x0GjGQvZ2q8l52/tczXl8VQyZl5KMmQaumTINHTJkGnQ5/kBxJ4CdtQ1jzkAAAAASUVORK5CYII="
)


def make_template(path: Path) -> None:
    doc = Document()

    particulars = doc.add_table(rows=6, cols=2)
    fields = [
        "ProjectName",
        "ProjectNumber",
        "InspectionDate",
        "InspectionTime",
        "Weather",
        "StageOfWorks",
    ]
    for row, field in zip(particulars.rows, fields):
        row.cells[0].text = field
        row.cells[1].text = "{{" + field + "}}"

    attendees = doc.add_table(rows=4, cols=2)
    attendees.cell(0, 0).text = "NAME"
    attendees.cell(0, 1).text = "COMPANY / ROLE"
    for row in attendees.rows[1:]:
        row.cells[0].text = "[name]"
        row.cells[1].text = "[role]"

    observations = doc.add_table(rows=3, cols=3)
    for index, label in enumerate(("ITEM", "DESCRIPTION", "ACTION / OWNER")):
        observations.cell(0, index).text = label
        observations.cell(1, index).text = "section"
        observations.cell(2, index).text = "item"

    doc.add_paragraph("PHOTOGRAPHS")
    doc.add_paragraph("Insert site photographs here")
    doc.add_paragraph("BDM Site Representative: __________________")
    doc.save(path)


class SiteInspectionPaginationTest(unittest.TestCase):
    def test_ten_portrait_photos_are_split_four_four_two(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            template = root / "331-test-template.docx"
            photos = root / "photos"
            photos.mkdir()
            make_template(template)

            for number in range(1, 11):
                # The builder intentionally discovers photo*.jpg names; python-docx
                # identifies the actual image format from the file header.
                (photos / f"photo{number}.jpg").write_bytes(PORTRAIT_IMAGE)

            config = {
                "template": str(template),
                "particulars": {
                    "ProjectName": "Synthetic pagination regression",
                    "ProjectNumber": "TEST",
                    "InspectionDate": "2026-08-13",
                    "InspectionTime": "7:15 AM",
                    "Weather": "Fine",
                    "StageOfWorks": "Structure",
                },
                "attendees": [["Test user", "BDM"]],
                "observations": [
                    ["1.0", "GENERAL", "", True],
                    ["1.1", "Synthetic observation referencing Photos 1-10.", "BDM", False],
                ],
                "photos_dir": str(photos),
                "captions": [
                    f"Photo {number} - Synthetic long caption for pagination regression testing (ref 1.1)."
                    for number in range(1, 11)
                ],
                "options": {"photos_per_page": 4, "signature": False},
            }
            config_path = root / "config.json"
            config_path.write_text(json.dumps(config), encoding="utf-8")
            output = root / "output.docx"

            result = subprocess.run(
                [sys.executable, str(BUILDER), str(config_path), str(output)],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

            report = Document(output)
            photo_tables = report.tables[3:]
            self.assertEqual([len(table.rows) for table in photo_tables], [2, 2, 1])
            self.assertEqual(len(report.inline_shapes), 10)

            continued = [
                paragraph
                for paragraph in report.paragraphs
                if "PHOTOGRAPHS (CONTINUED)" in paragraph.text
            ]
            self.assertEqual(len(continued), 2)
            self.assertTrue(
                all(paragraph._p.xpath('.//w:br[@w:type="page"]') for paragraph in continued)
            )

            captions = "\n".join(cell.text for table in photo_tables for row in table.rows for cell in row.cells)
            for number in range(1, 11):
                self.assertIn(f"Photo {number} -", captions)


if __name__ == "__main__":
    unittest.main()
